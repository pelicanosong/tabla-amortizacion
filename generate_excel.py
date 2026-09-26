import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import datetime
import shutil

def create_amortization_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Seguimiento y Pagos Reales"

    # Completamente desbloqueado / sin contraseñas
    ws.protection.sheet = False
    ws.protection.disable()

    font_family = "Calibri"
    
    # Paleta de colores profesionales
    c_navy = "0F172A"        # Slate 900
    c_accent_blue = "1E40AF" # Blue 800
    c_light_blue = "DBEAFE"  # Blue 100 para totales
    c_soft_gray = "F8FAFC"
    c_zebra = "F1F5F9"
    c_border = "CBD5E1"
    c_green_light = "DCFCE7" # Green 100
    c_input_bg = "FEF9C3"    # Amarillo suave editable
    c_real_pay_bg = "EFF6FF" # Azul suave columna pago real

    title_font = Font(name=font_family, size=15, bold=True, color="FFFFFF")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="93C5FD")
    section_font = Font(name=font_family, size=11, bold=True, color="1E3A8A")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    bold_font = Font(name=font_family, size=10, bold=True)
    normal_font = Font(name=font_family, size=10)
    note_font = Font(name=font_family, size=9, italic=True, color="475569")

    thin_border_side = Side(style='thin', color=c_border)
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom_side = Side(style='medium', color="0F172A")
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)

    title_fill = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    header_fill = PatternFill(start_color=c_accent_blue, end_color=c_accent_blue, fill_type="solid")
    input_fill = PatternFill(start_color=c_input_bg, end_color=c_input_bg, fill_type="solid")
    calc_fill = PatternFill(start_color=c_soft_gray, end_color=c_soft_gray, fill_type="solid")
    summary_fill = PatternFill(start_color=c_green_light, end_color=c_green_light, fill_type="solid")
    real_pay_fill = PatternFill(start_color=c_real_pay_bg, end_color=c_real_pay_bg, fill_type="solid")
    zebra_fill = PatternFill(start_color=c_zebra, end_color=c_zebra, fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Encabezado Principal
    ws.merge_cells("A1:I1")
    ws["A1"] = "CONTROL DE PRÉSTAMOS Y SEGUIMIENTO DE PAGOS REALES"
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    ws["A2"] = "Permite ingresar abonos reales variables (paga más, paga menos o no paga) con recálculo de saldo"
    ws["A2"].font = subtitle_font
    ws["A2"].fill = title_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Sección de Datos del Cliente y Préstamo
    ws.merge_cells("A4:D4")
    ws["A4"] = "DATOS DEL CLIENTE Y PRÉSTAMO"
    ws["A4"].font = section_font
    
    ws.merge_cells("F4:I4")
    ws["F4"] = "RESUMEN DE LIQUIDACIÓN Y ESTADO"
    ws["F4"].font = section_font

    inputs = [
        ("A5", "Nombre del Cliente:", "C5", "Carlos Alberto Rodríguez", '@'),
        ("A6", "Capital Inicial Prestado:", "C6", 50000000, '"$"#,##0.00'),
        ("A7", "Tasa Nominal Anual (TNA):", "C7", 0.12, '0.00%'),
        ("A8", "Tasa Mensual:", "C8", "=C7/12", '0.0000%'),
        ("A9", "Plazo Total (Meses):", "C9", 180, '#,##0'),
        ("A10", "Fecha de Inicio:", "C10", datetime.date(2026, 1, 1), 'DD/MM/YYYY'),
    ]

    for label_cell, label_text, val_cell, val, num_fmt in inputs:
        ws[label_cell] = label_text
        ws[label_cell].font = bold_font
        ws[label_cell].alignment = Alignment(horizontal="left", vertical="center")
        
        ws[val_cell] = val
        ws[val_cell].font = bold_font
        ws[val_cell].number_format = num_fmt
        ws[val_cell].alignment = Alignment(horizontal="right", vertical="center")
        ws[val_cell].border = thin_border
        
        if label_cell in ["A5", "A6", "A7", "A9", "A10"]:
            ws[val_cell].fill = input_fill
        else:
            ws[val_cell].fill = calc_fill

    # Resumen derecho
    summaries = [
        ("F5", "Cuota Fija Sugerida (Pactada):", "H5", '=IF(AND(C6>0, C9>0, C8>0), PMT(C8, C9, -C6), 0)', '"$"#,##0.00'),
        ("F6", "Saldo Pendiente Actual:", "H6", '=H194', '"$"#,##0.00'),
        ("F7", "Total Pagado Real a la Fecha:", "H7", '=SUM(F14:F193)', '"$"#,##0.00'),
        ("F8", "Total Intereses Pagados:", "H8", '=SUM(E14:E193)', '"$"#,##0.00'),
        ("F9", "Total Capital Amortizado:", "H9", '=SUM(G14:G193)', '"$"#,##0.00'),
        ("F10", "Fecha Final Pactada:", "H10", '=IF(C9>0, EDATE(C10, C9), "")', 'DD/MM/YYYY'),
    ]

    for label_cell, label_text, val_cell, formula, num_fmt in summaries:
        ws[label_cell] = label_text
        ws[label_cell].font = bold_font
        ws[label_cell].alignment = Alignment(horizontal="left", vertical="center")
        
        ws[val_cell] = formula
        ws[val_cell].font = bold_font
        ws[val_cell].number_format = num_fmt
        ws[val_cell].alignment = Alignment(horizontal="right", vertical="center")
        ws[val_cell].border = thin_border
        ws[val_cell].fill = summary_fill

    # Nota de uso
    ws.merge_cells("A11:I11")
    ws["A11"] = "💡 INSTRUCCIÓN: En la columna 'Pago Real Realizado' (Columna F) escribe exactamente lo que el cliente abonó ese mes. El abono a capital y el saldo se ajustan de inmediato."
    ws["A11"].font = note_font
    ws["A11"].alignment = Alignment(horizontal="left", vertical="center")

    # Títulos de las Columnas
    headers = [
        ("A13", "N° Cuota"),
        ("B13", "Fecha"),
        ("C13", "Capital Inicial"),
        ("D13", "Cuota Pactada"),
        ("E13", "Intereses del Mes"),
        ("F13", "Pago Real Realizado"),
        ("G13", "Abono a Capital"),
        ("H13", "Saldo Capital Real"),
        ("I13", "Estado del Pago")
    ]

    ws.row_dimensions[13].height = 26
    for pos, h_title in headers:
        ws[pos] = h_title
        ws[pos].font = header_font
        ws[pos].fill = header_fill
        ws[pos].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[pos].border = header_border

    # Generación de 180 filas (filas 14 a 193)
    total_months = 180
    start_row = 14

    for i in range(1, total_months + 1):
        row = start_row + i - 1
        ws.row_dimensions[row].height = 20
        fill_to_use = zebra_fill if i % 2 == 0 else white_fill

        # Col A: N°
        ws[f"A{row}"] = i
        ws[f"A{row}"].font = normal_font
        ws[f"A{row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"A{row}"].border = thin_border
        ws[f"A{row}"].fill = fill_to_use

        # Col B: Fecha
        ws[f"B{row}"] = f'=IF(A{row}<=$C$9, EDATE($C$10, A{row}), "")'
        ws[f"B{row}"].font = normal_font
        ws[f"B{row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"B{row}"].number_format = 'DD/MM/YYYY'
        ws[f"B{row}"].border = thin_border
        ws[f"B{row}"].fill = fill_to_use

        # Col C: Capital Inicial
        if i == 1:
            ws[f"C{row}"] = f'=IF(A{row}<=$C$9, $C$6, 0)'
        else:
            ws[f"C{row}"] = f'=IF(A{row}<=$C$9, H{row-1}, 0)'
        ws[f"C{row}"].font = normal_font
        ws[f"C{row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"C{row}"].number_format = '"$"#,##0.00'
        ws[f"C{row}"].border = thin_border
        ws[f"C{row}"].fill = fill_to_use

        # Col D: Cuota Pactada Sugerida
        ws[f"D{row}"] = f'=IF(AND(A{row}<=$C$9, C{row}>0), $H$5, 0)'
        ws[f"D{row}"].font = normal_font
        ws[f"D{row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"D{row}"].number_format = '"$"#,##0.00'
        ws[f"D{row}"].border = thin_border
        ws[f"D{row}"].fill = fill_to_use

        # Col E: Intereses del Mes (Capital Inicial * Tasa Mensual)
        ws[f"E{row}"] = f'=IF(AND(A{row}<=$C$9, C{row}>0), C{row}*$C$8, 0)'
        ws[f"E{row}"].font = normal_font
        ws[f"E{row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"E{row}"].number_format = '"$"#,##0.00'
        ws[f"E{row}"].border = thin_border
        ws[f"E{row}"].fill = fill_to_use

        # Col F: Pago Real Realizado (Default formula takes D{row}, but user can type custom value)
        # Sample payments for first 3 rows to demonstrate real cases:
        if i == 1:
            ws[f"F{row}"] = "=D14"  # Pago normal
        elif i == 2:
            ws[f"F{row}"] = 1500000  # Abono extraordinario ejemplo
        elif i == 3:
            ws[f"F{row}"] = 300000   # Pago parcial ejemplo
        else:
            ws[f"F{row}"] = f'=IF(AND(A{row}<=$C$9, C{row}>0), D{row}, 0)'
            
        ws[f"F{row}"].font = bold_font
        ws[f"F{row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"F{row}"].number_format = '"$"#,##0.00'
        ws[f"F{row}"].border = thin_border
        ws[f"F{row}"].fill = real_pay_fill

        # Col G: Abono a Capital (Pago Real - Intereses, limitado al saldo)
        ws[f"G{row}"] = f'=IF(AND(A{row}<=$C$9, C{row}>0), MIN(C{row}, MAX(0, F{row}-E{row})), 0)'
        ws[f"G{row}"].font = normal_font
        ws[f"G{row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"G{row}"].number_format = '"$"#,##0.00'
        ws[f"G{row}"].border = thin_border
        ws[f"G{row}"].fill = fill_to_use

        # Col H: Saldo Capital Real (Capital Inicial - Abono a Capital)
        ws[f"H{row}"] = f'=IF(AND(A{row}<=$C$9, C{row}>0), MAX(0, C{row}-G{row}), 0)'
        ws[f"H{row}"].font = bold_font
        ws[f"H{row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws[f"H{row}"].number_format = '"$"#,##0.00'
        ws[f"H{row}"].border = thin_border
        ws[f"H{row}"].fill = fill_to_use

        # Col I: Estado del Pago
        ws[f"I{row}"] = f'=IF(A{row}>$C$9, "", IF(C{row}=0, "Finalizado", IF(F{row}=0, "Sin Pago", IF(F{row}>D{row}+100, "Abono Extra", IF(F{row}<D{row}-100, "Pago Parcial", "Al Día")))))'
        ws[f"I{row}"].font = normal_font
        ws[f"I{row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"I{row}"].border = thin_border
        ws[f"I{row}"].fill = fill_to_use

    # Fila de Totales
    tot_row = start_row + total_months
    ws.row_dimensions[tot_row].height = 24
    ws[f"A{tot_row}"] = "TOTALES"
    ws[f"A{tot_row}"].font = bold_font
    ws[f"A{tot_row}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"A{tot_row}"].border = thin_border
    ws[f"A{tot_row}"].fill = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")

    for col in ["B", "C"]:
        ws[f"{col}{tot_row}"] = ""
        ws[f"{col}{tot_row}"].border = thin_border
        ws[f"{col}{tot_row}"].fill = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")

    ws[f"D{tot_row}"] = f'=SUM(D14:D{tot_row-1})'
    ws[f"D{tot_row}"].font = bold_font
    ws[f"D{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"D{tot_row}"].number_format = '"$"#,##0.00'
    ws[f"D{tot_row}"].border = thin_border
    ws[f"D{tot_row}"].fill = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")

    ws[f"E{tot_row}"] = f'=SUM(E14:E{tot_row-1})'
    ws[f"E{tot_row}"].font = bold_font
    ws[f"E{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"E{tot_row}"].number_format = '"$"#,##0.00'
    ws[f"E{tot_row}"].border = thin_border
    ws[f"E{tot_row}"].fill = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")

    ws[f"F{tot_row}"] = f'=SUM(F14:F{tot_row-1})'
    ws[f"F{tot_row}"].font = bold_font
    ws[f"F{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"F{tot_row}"].number_format = '"$"#,##0.00'
    ws[f"F{tot_row}"].border = thin_border
    ws[f"F{tot_row}"].fill = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")

    ws[f"G{tot_row}"] = f'=SUM(G14:G{tot_row-1})'
    ws[f"G{tot_row}"].font = bold_font
    ws[f"G{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"G{tot_row}"].number_format = '"$"#,##0.00'
    ws[f"G{tot_row}"].border = thin_border
    ws[f"G{tot_row}"].fill = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")

    ws[f"H{tot_row}"] = f'=H{tot_row-1}'
    ws[f"H{tot_row}"].font = bold_font
    ws[f"H{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"H{tot_row}"].number_format = '"$"#,##0.00'
    ws[f"H{tot_row}"].border = thin_border
    ws[f"H{tot_row}"].fill = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")

    ws[f"I{tot_row}"] = ""
    ws[f"I{tot_row}"].border = thin_border
    ws[f"I{tot_row}"].fill = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")

    # Anchos de columna
    col_widths = {
        "A": 12, # N°
        "B": 15, # Fecha
        "C": 20, # Capital Inicial
        "D": 18, # Cuota Pactada
        "E": 18, # Intereses
        "F": 22, # Pago Real Realizado
        "G": 18, # Abono a Capital
        "H": 20, # Saldo Capital Real
        "I": 16  # Estado
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    paths = [
        "/Users/nico/.gemini/antigravity/scratch/tabla_amortizacion/Tabla_Amortizacion_180_Meses.xlsx",
        "/Users/nico/Desktop/Tabla_Amortizacion_180_Meses.xlsx",
        "/Users/nico/Downloads/Tabla_Amortizacion_180_Meses.xlsx"
    ]
    
    wb.save(paths[0])
    shutil.copyfile(paths[0], paths[1])
    shutil.copyfile(paths[0], paths[2])
    print("Archivo Excel actualizado con seguimiento de pagos reales.")

if __name__ == "__main__":
    create_amortization_sheet()
